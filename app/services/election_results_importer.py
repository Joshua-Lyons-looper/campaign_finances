from pathlib import Path
import re

from openpyxl import load_workbook


class ElectionResultsImporter:
    def __init__(self, data_directory: str = "data"):
        self.data_directory = Path(data_directory)
    def list_sheets(self,cycle: int,) -> list[str]:
        file_path = (
            self.data_directory
            / f"federalelections{cycle}.xlsx"
        )
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )
        try:
            return workbook.sheetnames
        finally:
            workbook.close()

    def preview_sheet(self,cycle: int,sheet_name: str,rows: int = 20,) -> list[list]:
        file_path = (
            self.data_directory
            / f"federalelections{cycle}.xlsx"
        )
        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet = workbook[sheet_name]

            preview = []

            for row in worksheet.iter_rows(
                min_row=1,
                max_row=rows,
                min_col=1,
                max_col=30,
                values_only=True,
            ):
                preview.append(list(row))

            return preview

        finally:
            workbook.close()

    def read_election_results(self,cycle: int,office: str,) -> list[dict]:
        office = office.upper()
        if office == "H":
            sheet_suffix = "US House Results by State"
        elif office == "S":
            sheet_suffix = "US Senate Results by State"
        elif office == "P":
            sheet_suffix = "Pres General Results"
        else:
            raise ValueError("office must be P, H, or S")

        
        file_path = (
            self.data_directory
            / f"federalelections{cycle}.xlsx"
        )

        workbook = load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet_name = next(
                (
                    name
                    for name in workbook.sheetnames
                    if name.endswith(sheet_suffix)
                ),
                None,
            )

            combined_sheet = False
            if sheet_name is None and office in {"H", "S"}:
                sheet_name = next(
                    (
                        name for name in workbook.sheetnames
                        if re.search(
                            r"\bUS\s+House\s*&\s*Senate\s+Resu(?:l)?ts\s*$",
                            name,
                            re.IGNORECASE,
                        )
                    ),
                    None,
                )
                combined_sheet = sheet_name is not None

            if sheet_name is None:
                raise ValueError(
                    f"Could not find '{sheet_suffix}' "
                    f"in the {cycle} election workbook. "
                    f"Available sheets: {workbook.sheetnames!r}"
                )
            worksheet = workbook[sheet_name]

            if office in {"H", "S"}:
                headers = next(worksheet.iter_rows(max_row=1, values_only=True))
                columns = {
                    header.strip().upper(): index
                    for index, header in enumerate(headers)
                    if isinstance(header, str)
                }
                # Older combined sheets put the state name before its abbreviation
                # and may omit the general-election winner indicator entirely.
                state_column = columns.get("STATE ABBREVIATION", 1)
                winner_column = columns.get("GE WINNER INDICATOR")

            results = []

            for row in worksheet.iter_rows(min_row=2,values_only=True,):
                if office == "P":
                    # Presidential worksheet has a different column layout
                    state = (
                        row[3].strip()
                        if isinstance(row[3], str)
                        else None
                    )

                    fec_candidate_id = (
                        row[1].strip()
                        if isinstance(row[1], str)
                        else None
                    )

                    candidate_name = (
                        row[7].strip()
                        if isinstance(row[7], str)
                        else None
                    )

                    party = (
                        row[9].strip()
                        if isinstance(row[9], str)
                        else None
                    )

                    general_votes = row[10]
                    general_percentage = row[11]

                    general_winner = (
                        row[15].strip()
                        if isinstance(row[15], str)
                        else row[15]
                    )

                    district = None

                else:
                    # House and Senate worksheet layout
                    state = (
                        row[state_column].strip()
                        if isinstance(row[state_column], str)
                        else None
                    )

                    fec_candidate_id = (
                        row[4].strip()
                        if isinstance(row[4], str)
                        else None
                    )

                    # Candidate ID prefixes distinguish offices on combined sheets.
                    if combined_sheet and (
                        not fec_candidate_id
                        or not fec_candidate_id.startswith(office)
                    ):
                        continue

                    candidate_name = (
                        row[8].strip()
                        if isinstance(row[8], str)
                        else None
                    )

                    party = (
                        row[10].strip()
                        if isinstance(row[10], str)
                        else None
                    )

                    general_votes = row[15]
                    general_percentage = row[16]

                    general_winner = None
                    if winner_column is not None:
                        general_winner = (
                            row[winner_column].strip()
                            if isinstance(row[winner_column], str)
                            else row[winner_column]
                        )

                    if office == "H":
                        district = (
                            str(row[3]).strip().zfill(2)
                            if row[3] is not None
                            else None
                        )
                    else:
                        district = None

                if not fec_candidate_id:
                    continue

                if fec_candidate_id.lower() == "n/a":
                    continue

                if general_votes is None:
                    continue

                results.append({
                    "fec_candidate_id": fec_candidate_id,
                    "cycle": cycle,
                    "state": state,
                    "office": office,
                    "district": district,
                    "candidate_name": candidate_name,
                    "party": party,
                    "general_votes": general_votes,
                    "general_percentage": general_percentage,
                    "general_winner": general_winner,
                })
            return results

        finally:
            workbook.close()
